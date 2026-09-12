"""Validate source workbooks and expose explicit, reproducible time conventions.

Main arrays use right-endpoint power as the representative 10-minute power.
The *_trapezoid arrays provide the alternative piecewise-linear integral.
No fitted predictor, future-observation imputation, or outlier deletion is used.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, time, timedelta
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
EXPECTED_DATES = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
RELEASE_HOURS = (0, 6, 12, 18)


def minutes(value: object) -> int:
    """Convert mixed Excel time cells, preserving the explicit next-day mark."""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"Unexpected fractional time: {value!r}")
        return value.hour * 60 + value.minute
    if isinstance(value, str):
        text = value.strip()
        day_offset = 1440 if text.endswith("+1") else 0
        if day_offset:
            text = text[:-2]
        hour, minute = map(int, text.split(":"))
        return day_offset + hour * 60 + minute
    raise ValueError(f"Unsupported time cell: {value!r}")


def calendar_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return date(*map(int, value.split("-")))
    raise ValueError(f"Unsupported date: {value!r}")


def checked_numeric(rows: list, shape: tuple[int, ...], label: str) -> np.ndarray:
    # float(None) fails rather than silently treating a blank as zero.
    out = np.asarray([[float(v) for v in row] for row in rows], dtype=np.float64)
    if out.shape != shape:
        raise ValueError(f"{label}: expected {shape}, received {out.shape}")
    if not np.all(np.isfinite(out)):
        raise ValueError(f"{label}: contains non-finite values")
    if np.any(out < 0):
        raise ValueError(f"{label}: contains negative values; inspect source")
    return out


def read_daily(sheet, label: str) -> np.ndarray:
    rows = list(sheet.values)
    if [minutes(v) for v in rows[0][1:]] != list(range(10, 1441, 10)):
        raise ValueError(f"{label}: time columns must run from 00:10 through 24:00")
    if [calendar_date(row[0]) for row in rows[1:]] != EXPECTED_DATES:
        raise ValueError(f"{label}: dates must be every day of 2025 in order")
    return checked_numeric([row[1:] for row in rows[1:]], (365, 144), label)


def trapezoid_daily(values: np.ndarray) -> np.ndarray:
    """Interval means for a continuous annual piecewise-linear power curve.

    The one missing annual origin is held equal to the first sample. Subsequent
    days start at the immediately preceding day's 24:00 observation.
    """
    flat = values.ravel()
    left = np.concatenate(([flat[0]], flat[:-1]))
    return ((left + flat) / 2).reshape(values.shape)


def summarize(values: np.ndarray) -> dict:
    return {
        "shape": list(values.shape),
        "minimum": float(np.min(values)),
        "maximum": float(np.max(values)),
        "mean": float(np.mean(values)),
        "missing_or_nonfinite": int(np.count_nonzero(~np.isfinite(values))),
        "negative_count": int(np.count_nonzero(values < 0)),
    }


def prepare(raw: Path, output: Path) -> dict:
    paths = {i: raw / f"附件{i}.xlsx" for i in range(1, 5)}
    with paths[1].open("rb") as stream:
        wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            rows = list(wb["Sheet1"].values)
        finally:
            wb.close()
    if [minutes(row[0]) for row in rows[1:]] != list(range(10, 1441, 10)):
        raise ValueError("Attachment 1: invalid time sequence")
    q1 = checked_numeric([row[1:] for row in rows[1:]], (144, 3), "Attachment 1")

    wb = openpyxl.load_workbook(paths[2], read_only=True, data_only=True)
    try:
        load = read_daily(wb["小区负载"], "Load")
        pv = read_daily(wb["光伏发电实际功率"], "Actual PV")
    finally:
        wb.close()
    wb = openpyxl.load_workbook(paths[4], read_only=True, data_only=True)
    try:
        price = read_daily(wb["Sheet1"], "Dynamic price")
    finally:
        wb.close()
    wb = openpyxl.load_workbook(paths[3], read_only=True, data_only=True)
    try:
        forecast_rows = list(wb["Sheet1"].values)
    finally:
        wb.close()
    if list(forecast_rows[0][2:]) != [f"预报{i}小时" for i in range(1, 25)]:
        raise ValueError("Attachment 3: forecast horizons are not hours +1 through +24")
    if len(forecast_rows) != 1461:
        raise ValueError("Attachment 3: expected 1460 issue rows")
    inherited = None
    for i, row in enumerate(forecast_rows[1:]):
        if row[0] not in (None, ""):
            inherited = calendar_date(row[0])
        expected_date = EXPECTED_DATES[i // 4]
        expected_hour = RELEASE_HOURS[i % 4]
        if inherited != expected_date or minutes(row[1]) != expected_hour * 60:
            raise ValueError(f"Attachment 3: incorrect issue date/hour at row {i + 2}")
    forecast_hourly = checked_numeric(
        [row[2:] for row in forecast_rows[1:]], (1460, 24), "PV forecasts"
    ).reshape(365, 4, 24)

    # The interpolation origin is available at issue time, not after it.
    anchors = np.zeros((365, 4), dtype=np.float64)
    anchors[1:, 0] = pv[:-1, -1]
    for issue_index, hour in enumerate(RELEASE_HOURS[1:], start=1):
        anchors[:, issue_index] = pv[:, hour * 6 - 1]
    hourly_knots = np.concatenate((anchors[..., None], forecast_hourly), axis=2)
    fine_knots = np.empty((365, 4, 145), dtype=np.float64)
    for d in range(365):
        for issue_index in range(4):
            fine_knots[d, issue_index] = np.interp(
                np.arange(145) / 6, np.arange(25), hourly_knots[d, issue_index]
            )

    arrays = {
        "dates": np.array([d.isoformat() for d in EXPECTED_DATES]),
        "load": load,
        "pv": pv,
        "fixed_price": q1[:, 0],
        "dynamic_price": price,
        "q1_load": q1[:, 1],
        "q1_pv": q1[:, 2],
        "forecast": fine_knots[:, :, 1:],
        "load_trapezoid": trapezoid_daily(load),
        "pv_trapezoid": trapezoid_daily(pv),
        "q1_load_trapezoid": (np.roll(q1[:, 1], 1) + q1[:, 1]) / 2,
        "q1_pv_trapezoid": (np.roll(q1[:, 2], 1) + q1[:, 2]) / 2,
        "forecast_trapezoid": (fine_knots[:, :, :-1] + fine_knots[:, :, 1:]) / 2,
        "forecast_hourly": forecast_hourly,
        "forecast_anchor": anchors,
        "forecast_release_hours": np.array(RELEASE_HOURS, dtype=np.int64),
        "interval_end_minutes": np.arange(10, 1441, 10, dtype=np.int64),
        "dt_hours": np.array(1 / 6),
    }
    # At exact hourly targets, the right-endpoint interpolant must reproduce
    # the supplied forecasts; no day shift or averaging is allowed here.
    np.testing.assert_allclose(arrays["forecast"][:, :, 5::6], forecast_hourly,
                               rtol=0, atol=1e-10)
    # Integration of the linear interpolant must preserve hourly trapezoids.
    np.testing.assert_allclose(
        arrays["forecast_trapezoid"].reshape(365, 4, 24, 6).mean(axis=3),
        (hourly_knots[:, :, :-1] + hourly_knots[:, :, 1:]) / 2,
        rtol=1e-13, atol=1e-10,
    )
    output.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(output / "data.npz", **arrays)
    audit = {
        "source_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                          for p in paths.values()},
        "date_start": "2025-01-01",
        "date_end": "2025-12-31",
        "days": 365,
        "submission_start": "2025-02-01",
        "submission_days": 334,
        "interval_minutes": 10,
        "main_power_convention": "right_endpoint_representative_interval_mean",
        "forecast_origin": "PV observation at issue time; Jan 1 00:00 = 0 kW",
        "statistics": {key: summarize(value) for key, value in arrays.items()
                       if key != "dates"},
        "annual_energy_kwh": {"load": float(load.sum() / 6),
                              "pv": float(pv.sum() / 6),
                              "load_trapezoid": float(arrays["load_trapezoid"].sum() / 6),
                              "pv_trapezoid": float(arrays["pv_trapezoid"].sum() / 6)},
        "q1_energy_kwh": {"load": float(q1[:, 1].sum() / 6),
                          "pv": float(q1[:, 2].sum() / 6)},
        "a1_versus_annual_same_time_average_max_abs_kw": {
            "load": float(np.max(np.abs(load.mean(axis=0) - q1[:, 1]))),
            "pv": float(np.max(np.abs(pv.mean(axis=0) - q1[:, 2]))),
        },
        "duplicate_daily_rows": {"load": 365 - len(np.unique(load, axis=0)),
                                 "pv": 365 - len(np.unique(pv, axis=0)),
                                 "dynamic_price": 365 - len(np.unique(price, axis=0))},
        "versions": {"numpy": np.__version__, "openpyxl": openpyxl.__version__},
    }
    (output / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
                                       encoding="utf-8")
    return audit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", type=Path, default=ROOT / "data" / "raw")
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "processed")
    args = parser.parse_args()
    audit = prepare(args.raw, args.output)
    print(json.dumps({"output": str(args.output / "data.npz"),
                      "days": audit["days"],
                      "annual_energy_kwh": audit["annual_energy_kwh"]},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
